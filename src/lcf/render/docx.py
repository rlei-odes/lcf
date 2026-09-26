"""The docx deliverable.

Two paths, because a rule builder should not have to produce a Word template
before anyone can get a document out, and should be able to when branding matters.

- `render_plain` builds a clean document directly with python-docx. Works for any
  document type with no setup, real Word tables, real heading styles.
- `render_with_template` merges the same content into a rule builder's `.docx`,
  where the logo, fonts, header and CI colours already live (DESIGN §8). The
  template holds the design; this holds none of it.

`starter` closes the gap between them: it writes the template the second path
wants, from the spec, onto the house style — every tag already spelled correctly.
A tag nobody typed cannot be misspelled, which is the same bet the structured spec
editor makes (ARCHITECTURE §15.5).

The renderer stays dumb on purpose: no business rules, no model, no conditional
content beyond what the template itself expresses.
"""

import io
import re
from dataclasses import dataclass
from typing import Any

from docx import Document as NewDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docxtpl import DocxTemplate, RichText
from markdown_it import MarkdownIt

from lcf.engine.view import DocumentView
from lcf.spec.models import Block, BlockKind, DocTypeSpec, Section

# What a template may say. `{{ a.b }}`, `{{r a.b }}`, and the statement forms
# docxtpl adds for repeating a table row (`{%tr %}`) or a paragraph (`{%p %}`).
# The `r` of docxtpl's RichText form must be followed by space, or this eats the
# `r` of `{{ row.name }}` and lints a section called `ow`.
EXPR = re.compile(r"{{[-\s]*(?:r\s)?\s*(.+?)[-\s]*}}", re.DOTALL)
STMT = re.compile(r"{%[a-z]*[-\s]+(.+?)[-\s]*%}", re.DOTALL)
FOR = re.compile(r"^for\s+(\w+)\s+in\s+(\S+)")
PATH = re.compile(r"\b([a-zA-Z_]\w*)\.([a-zA-Z_]\w*)")

# Names that are not sections. `sections` is the whole ordered list, for a
# template that would rather loop than name each one.
ROOT = {"title", "doc_type", "sections"}

# The attributes a block exposes, by kind. Anything else named after a block is a
# template asking for something that will never arrive.
ATTRIBUTES: dict[BlockKind, set[str]] = {
    BlockKind.PROSE: {"text", "filled", "rich", "paragraphs"},
    BlockKind.LIST: {"text", "filled", "items"},
    BlockKind.KEYVALUE: {"text", "filled", "rows", "columns"},
    BlockKind.TABLE: {"text", "filled", "rows", "columns"},
    BlockKind.IMAGE_REF: {"text", "filled", "rows"},
}
SECTION_ATTRIBUTES = {"title", "key"}


class BlockValue(str):
    """A block's flat text, with its structure still reachable.

    `{{ d1_team.members }}` gives the flattened string it always gave, so a
    template written against the first version of this keeps working. `.rows`,
    `.items` and `.paragraphs` give the same content unflattened, which is what a
    branded template actually wants: a real Word table row repeated by `{%tr %}`
    keeps the designer's table style, and a flattened string cannot.
    """

    filled: bool
    rows: list[dict[str, Any]]
    columns: list[dict[str, str]]
    items: list[str]
    paragraphs: list[RichText]
    rich: RichText

    def __new__(cls, text: str, **attrs: Any) -> "BlockValue":
        value = super().__new__(cls, text)
        value.filled = bool(text)
        value.rows = attrs.get("rows", [])
        value.columns = attrs.get("columns", [])
        value.items = attrs.get("items", [])
        value.paragraphs = attrs.get("paragraphs", [])
        value.rich = attrs.get("rich") or RichText(text)
        return value

    @property
    def text(self) -> str:
        return str(self)


@dataclass
class TemplateLint:
    """What an upload has to answer: is it wrong, and is it complete.

    Separate lists because they are different severities. A tag naming a section
    that does not exist renders as nothing and is always a mistake. A section the
    template never mentions may be deliberate — an internal section the customer
    does not see — so it is said out loud and allowed.
    """

    problems: list[str]
    missing: list[str]

    @property
    def ok(self) -> bool:
        return not self.problems

    @property
    def complete(self) -> bool:
        return not self.missing


def render_plain(view: DocumentView, *, title: str, base: bytes | None = None) -> bytes:
    """A presentable document with no template at all.

    `base` is the house style — header, footer, logo, fonts. Given one, this is
    still the plain renderer, just not on blank paper.
    """
    doc = _open(base)
    doc.add_heading(title, level=0)

    for section in view.spec.sections:
        filled = [
            (block, view.block_value(section.key, block.key))
            for block in section.blocks
            if not _blank(view.block_value(section.key, block.key))
        ]
        if not filled:
            continue

        doc.add_heading(section.title, level=1)
        for block, value in filled:
            if not _repeats_heading(section, block):
                doc.add_heading(block.label, level=2)
            _write_block(doc, block, value)

    return _save(doc)


def _repeats_heading(section, block) -> bool:
    """Whether this block's label would only say the section heading again.

    The common shape of a section is one paragraph, and the natural name for that
    paragraph is the section's own — which reads well in the editor and as a
    stuttering pair of headings in the export. One block, same name: the section
    heading has already said it.
    """
    return len(section.blocks) == 1 and block.label.strip().lower() == section.title.strip().lower()


def render_with_template(view: DocumentView, template: bytes, *, title: str) -> bytes:
    """Merge content into a rule builder's branded template."""
    tpl = DocxTemplate(io.BytesIO(template))
    tpl.render(context(view, title=title, tpl=tpl))
    buffer = io.BytesIO()
    tpl.save(buffer)
    return buffer.getvalue()


def context(
    view: DocumentView, *, title: str, tpl: DocxTemplate | None = None
) -> dict[str, Any]:
    """What a template can reference: `{{ section_key.block_key }}`.

    Values are flattened to text here rather than in the template, so the simplest
    possible template cannot get the shape wrong — and carry their structure as
    attributes, so a template that wants a real table can have one.

    `tpl` is needed only to register hyperlink targets; without it a link in prose
    renders as its text, which is what a caller outside `render_with_template`
    (a test, the starter preview) wants anyway.
    """
    data: dict[str, Any] = {"title": title, "doc_type": view.spec.id}
    ordered: list[dict[str, Any]] = []
    for section in view.spec.sections:
        entry: dict[str, Any] = {"title": section.title, "key": section.key}
        for block in section.blocks:
            entry[block.key] = _value(block, view.block_value(section.key, block.key), tpl)
        data[section.key] = entry
        ordered.append(entry)
    data["sections"] = ordered
    return data


def _value(block: Block, value: Any, tpl: DocxTemplate | None) -> BlockValue:
    """One block, as both text and structure."""
    text = _as_text(block, value)
    if _blank(value):
        return BlockValue("")

    if block.kind is BlockKind.PROSE:
        paragraphs = _rich_paragraphs(str(value), tpl)
        return BlockValue(
            text, paragraphs=paragraphs, rich=_rich_paragraphs(str(value), tpl, joined=True)
        )

    if block.kind is BlockKind.LIST:
        return BlockValue(text, items=[str(item) for item in value])

    if block.kind is BlockKind.KEYVALUE:
        rows = [
            {"key": f.key, "label": f.label, "value": _cell(value.get(f.key))}
            for f in block.fields
            if not _blank(value.get(f.key))
        ]
        columns = [{"key": "label", "label": ""}, {"key": "value", "label": ""}]
        return BlockValue(text, rows=rows, columns=columns)

    if block.kind is BlockKind.TABLE:
        rows = [{c.key: _cell(row.get(c.key)) for c in block.columns} for row in value]
        columns = [{"key": c.key, "label": c.label} for c in block.columns]
        return BlockValue(text, rows=rows, columns=columns)

    if block.kind is BlockKind.IMAGE_REF:
        rows = [
            ref if isinstance(ref, dict) else {"caption": str(ref)}
            for ref in value
        ]
        return BlockValue(text, rows=rows)

    return BlockValue(text)


# --------------------------------------------------------------------------- #
# markdown → RichText
# --------------------------------------------------------------------------- #

_md = MarkdownIt("commonmark")


def _rich_paragraphs(
    source: str, tpl: DocxTemplate | None, *, joined: bool = False
) -> list[RichText] | RichText:
    """Prose is markdown (DESIGN §9); `{{ }}` would print its asterisks.

    Walks the token stream rather than pattern-matching the text — the subset is
    small, but a regex over markdown is the kind of thing that is wrong for a year
    before anyone notices. When `markdown/` lands as its own module this should
    consume its normalised AST instead of parsing here.
    """
    blocks: list[RichText] = []
    current = RichText()
    bold = italic = 0
    url_id: str | None = None
    prefix = ""

    for token in _md.parse(source or ""):
        if token.type == "bullet_list_open":
            prefix = "• "
        elif token.type == "ordered_list_open":
            prefix = "- "
        elif token.type in ("bullet_list_close", "ordered_list_close"):
            prefix = ""
        elif token.type == "inline":
            for child in token.children or []:
                if child.type == "text":
                    current.add(child.content, bold=bool(bold), italic=bool(italic), url_id=url_id)
                elif child.type == "code_inline":
                    current.add(child.content, font="Consolas", url_id=url_id)
                elif child.type == "strong_open":
                    bold += 1
                elif child.type == "strong_close":
                    bold = max(0, bold - 1)
                elif child.type == "em_open":
                    italic += 1
                elif child.type == "em_close":
                    italic = max(0, italic - 1)
                elif child.type == "link_open" and tpl is not None:
                    url_id = tpl.build_url_id(child.attrGet("href") or "")
                elif child.type == "link_close":
                    url_id = None
                elif child.type in ("softbreak", "hardbreak"):
                    current.add("\n")
        elif token.type in ("paragraph_close", "heading_close"):
            blocks.append(current)
            current = RichText()
        elif token.type == "list_item_close":
            blocks.append(current)
            current = RichText()
        elif token.type == "list_item_open" and prefix:
            current.add(prefix)

    if str(current.xml):
        blocks.append(current)

    if not joined:
        return blocks

    whole = RichText()
    for index, part in enumerate(blocks):
        if index:
            whole.add("\n")
        whole.xml += part.xml
    return whole


# --------------------------------------------------------------------------- #
# linting a template against the spec
# --------------------------------------------------------------------------- #


def lint(spec: DocTypeSpec, template: bytes) -> TemplateLint:
    """Does this template name things that exist, and does it name them all.

    Checked when a template is uploaded, so a template referring to a section that
    was renamed fails then — not at export time, in front of a customer.
    """
    text = _template_text(DocxTemplate(io.BytesIO(template)))
    problems: list[str] = []
    referenced: set[str] = set()
    blocks_seen: set[tuple[str, str]] = set()
    loop_vars: set[str] = set()

    # Loop variables first: `{{ row.name }}` inside a `{%tr for row in … %}` is
    # not a section reference and must not be linted as one.
    statements = STMT.findall(text)
    for statement in statements:
        match = FOR.match(statement.strip())
        if match:
            loop_vars.add(match.group(1))

    for statement in statements:
        body = statement.strip()
        if body.startswith(("end", "else", "elif ")) and not body.startswith("elif "):
            continue
        match = FOR.match(body)
        paths = [match.group(2)] if match else [f"{a}.{b}" for a, b in PATH.findall(body)]
        for path in paths:
            _check(spec, path, loop_vars, problems, referenced, blocks_seen)

    for expression in EXPR.findall(text):
        _check(spec, expression.strip(), loop_vars, problems, referenced, blocks_seen)

    return TemplateLint(sorted(set(problems)), _missing(spec, referenced, blocks_seen))


def _check(
    spec: DocTypeSpec,
    path: str,
    loop_vars: set[str],
    problems: list[str],
    referenced: set[str],
    blocks_seen: set[tuple[str, str]],
) -> None:
    parts = [p for p in re.split(r"[.\[]", path.split("|")[0].strip()) if p]
    if not parts:
        return
    head = parts[0]
    if head in loop_vars or head in ROOT or not re.fullmatch(r"[a-zA-Z_]\w*", head):
        return

    tag = f"{{{{ {path} }}}}"
    section = spec.section(head)
    if section is None:
        problems.append(f"{tag}: no section {head!r}")
        return

    referenced.add(head)
    if len(parts) == 1:
        return

    name = parts[1]
    if name in SECTION_ATTRIBUTES:
        return

    block = section.block(name)
    if block is None:
        problems.append(f"{tag}: {head!r} has no block {name!r}")
        return

    blocks_seen.add((head, name))
    if len(parts) > 2 and parts[2] not in ATTRIBUTES[block.kind]:
        allowed = ", ".join(sorted(ATTRIBUTES[block.kind]))
        problems.append(
            f"{tag}: {name!r} is a {block.kind.value} block and has no {parts[2]!r}. "
            f"Try one of: {allowed}"
        )


def _missing(
    spec: DocTypeSpec, referenced: set[str], blocks_seen: set[tuple[str, str]]
) -> list[str]:
    """Sections and blocks the template never mentions.

    The check that matters once a type can grow: a rule builder adds a section in
    the spec editor and the template, written against the version before it, is
    silently one chapter short.
    """
    notes: list[str] = []
    for section in spec.sections:
        if section.key not in referenced:
            notes.append(f"{section.title} ({section.key}) is not in the template")
            continue
        for block in section.blocks:
            if (section.key, block.key) not in blocks_seen:
                notes.append(f"{section.title}: {block.label} ({block.key}) is not in the template")
    return notes


def lint_template(spec: DocTypeSpec, template: bytes) -> list[str]:
    """The blocking half of `lint`, kept as its own name because it is the
    question an upload asks first."""
    return lint(spec, template).problems


def _template_text(tpl: DocxTemplate) -> str:
    # `.docx` stays None until the template is opened; get_docx() does that.
    doc = tpl.get_docx()
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    for section in doc.sections:
        for container in (section.header, section.footer):
            parts.extend(p.text for p in container.paragraphs)
    return "\n".join(parts)


# --------------------------------------------------------------------------- #
# generating the starter template
# --------------------------------------------------------------------------- #


def starter(spec: DocTypeSpec, *, base: bytes | None = None) -> bytes:
    """The template a rule builder edits instead of writes.

    Every section and block of the spec, as headings and correctly spelled tags,
    appended to the house style so the logo, fonts and header are already there.
    The author opens it in Word, moves things, deletes what the customer must not
    see, and uploads it back. What they never have to do is type a section key.
    """
    doc = _open(base)
    doc.add_heading("{{ title }}", level=0)

    for section in spec.sections:
        doc.add_heading(section.title, level=1)
        for block in section.blocks:
            if not _repeats_heading(section, block):
                doc.add_heading(block.label, level=2)
            _write_tags(doc, section, block)

    return _save(doc)


def _write_tags(doc, section: Section, block: Block) -> None:
    ref = f"{section.key}.{block.key}"

    if block.kind is BlockKind.PROSE:
        # Three paragraphs, the middle one repeated: `{%p %}` consumes the
        # paragraph it sits in, so the loop markers cannot be the content.
        doc.add_paragraph("{%p for para in " + ref + ".paragraphs %}")
        doc.add_paragraph("{{r para }}")
        doc.add_paragraph("{%p endfor %}")
        return

    if block.kind is BlockKind.LIST:
        doc.add_paragraph("{%p for item in " + ref + ".items %}")
        doc.add_paragraph("{{ item }}", style="List Bullet")
        doc.add_paragraph("{%p endfor %}")
        return

    if block.kind is BlockKind.KEYVALUE:
        _tag_table(doc, ref, [("label", ""), ("value", "")], header=False)
        return

    if block.kind is BlockKind.TABLE:
        _tag_table(doc, ref, [(c.key, c.label) for c in block.columns], header=True)
        return

    if block.kind is BlockKind.IMAGE_REF:
        doc.add_paragraph("{%p for image in " + ref + ".rows %}")
        doc.add_paragraph("{{ image.caption }}")
        doc.add_paragraph("{%p endfor %}")
        return

    doc.add_paragraph("{{ " + ref + " }}")


def _tag_table(doc, ref: str, columns: list[tuple[str, str]], *, header: bool) -> None:
    """A Word table whose body row repeats, so the author's table style survives.

    Same shape as the paragraph loop: a control row, the row that repeats, and a
    closing control row. Putting the tags in their own rows is the form docxtpl
    is unambiguous about, and an author who drags the table into their own style
    keeps the loop intact.
    """
    table = doc.add_table(rows=0, cols=len(columns))
    table.style = "Table Grid"

    if header:
        cells = table.add_row().cells
        for cell, (_, label) in zip(cells, columns, strict=True):
            cell.text = label
            _bold(cell)

    table.add_row().cells[0].text = "{%tr for row in " + ref + ".rows %}"
    cells = table.add_row().cells
    for cell, (key, _) in zip(cells, columns, strict=True):
        cell.text = "{{ row." + key + " }}"
    table.add_row().cells[0].text = "{%tr endfor %}"


# --------------------------------------------------------------------------- #
# writing blocks into a real document
# --------------------------------------------------------------------------- #


def _open(base: bytes | None):
    """A blank document, or the house style with its body emptied.

    A company template carries its design in the header, footer, styles and
    section properties — none of which live in the body. Whatever text is in the
    body is placeholder prose from whoever made it, and keeping it would put
    "Lorem ipsum" above every export.
    """
    if base is None:
        return NewDocument()
    doc = NewDocument(io.BytesIO(base))
    body = doc.element.body
    for child in list(body):
        if child.tag.endswith("}sectPr"):
            continue
        body.remove(child)
    return doc


def _save(doc) -> bytes:
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _write_block(doc, block: Block, value: Any) -> None:
    if block.kind is BlockKind.PROSE:
        for paragraph in str(value).split("\n\n"):
            if paragraph.strip():
                doc.add_paragraph(" ".join(paragraph.split()))
        return

    if block.kind is BlockKind.LIST:
        for item in value:
            doc.add_paragraph(str(item), style="List Bullet")
        return

    if block.kind is BlockKind.KEYVALUE:
        rows = [(f.label, value.get(f.key)) for f in block.fields if not _blank(value.get(f.key))]
        if not rows:
            return
        table = doc.add_table(rows=0, cols=2)
        table.style = "Light Grid Accent 1"
        for label, cell_value in rows:
            cells = table.add_row().cells
            cells[0].text = label
            cells[1].text = _cell(cell_value)
            _bold(cells[0])
        return

    if block.kind is BlockKind.TABLE:
        if not value:
            return
        table = doc.add_table(rows=1, cols=len(block.columns))
        table.style = "Light Grid Accent 1"
        for cell, column in zip(table.rows[0].cells, block.columns, strict=True):
            cell.text = column.label
            _bold(cell)
        for row in value:
            cells = table.add_row().cells
            for cell, column in zip(cells, block.columns, strict=True):
                cell.text = _cell(row.get(column.key))
        return

    if block.kind is BlockKind.IMAGE_REF:
        for ref in value:
            caption = ref.get("caption", "") if isinstance(ref, dict) else str(ref)
            paragraph = doc.add_paragraph(f"[image] {caption}")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        return

    doc.add_paragraph(str(value))


def _bold(cell) -> None:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.bold = True
            run.font.size = Pt(9)


def _as_text(block: Block, value: Any) -> str:
    """Flatten a block for a template placeholder."""
    if _blank(value):
        return ""
    if block.kind is BlockKind.PROSE:
        return str(value).strip()
    if block.kind is BlockKind.LIST:
        return "\n".join(f"• {item}" for item in value)
    if block.kind is BlockKind.KEYVALUE:
        return "\n".join(
            f"{f.label}: {_cell(value.get(f.key))}"
            for f in block.fields
            if not _blank(value.get(f.key))
        )
    if block.kind is BlockKind.TABLE:
        return "\n".join(
            " | ".join(f"{c.label}: {_cell(row.get(c.key))}" for c in block.columns)
            for row in value
        )
    if block.kind is BlockKind.IMAGE_REF:
        return "\n".join(
            str(ref.get("caption", "")) if isinstance(ref, dict) else str(ref) for ref in value
        )
    return str(value)


def _cell(value: Any) -> str:
    if value is None or value == "":
        return ""
    if value is True:
        return "yes"
    if value is False:
        return "no"
    return " ".join(str(value).split())


def _blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, bool | int | float):
        return False
    if isinstance(value, str):
        return not value.strip()
    return len(value) == 0
