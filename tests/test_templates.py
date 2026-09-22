"""The docx template, from upload to export.

The pure half — generating a starter, linting it, rendering through it — lives in
test_render.py and needs nothing. This is the half with a database and a bucket
behind it: that a template is bound to a version, carried forward to the next
one, and reported as incomplete when the spec grew past it.
"""

import io
import uuid

import pytest
from docx import Document as ReadDocx
from fastapi.testclient import TestClient
from sqlalchemy import delete, select

from lcf.core.db import session
from lcf.models.tables import DocType, DocTypeVersion, Document
from lcf.render import docx as docx_render
from lcf.services import doc_types
from lcf.services import templates as templates_service
from lcf.spec.models import Block, BlockKind, Section
from lcf.web.app import app


@pytest.fixture
def client():
    with TestClient(app) as c:
        yield c


@pytest.fixture
def storage():
    """Skip rather than fail when the bucket from .env is unreachable.

    Same bargain as the database fixture: the parts that must always run are the
    ones that need nothing.
    """
    from lcf.core.config import settings

    if not settings().s3_endpoint:
        pytest.skip("object storage not configured")
    try:
        from lcf.storage.s3 import client as s3

        s3().list_buckets()
    except Exception as exc:  # noqa: BLE001 — any failure to reach it is a skip
        pytest.skip(f"object storage not reachable: {exc}")


@pytest.fixture
async def published(db, spec_4d):
    """A uniquely-keyed 4D, removed again afterwards."""
    spec_4d.id = f"test-tpl-{uuid.uuid4().hex[:8]}"
    async with session() as s:
        await doc_types.publish(s, spec_4d)
    yield spec_4d
    async with session() as s:
        doc_type = await s.scalar(select(DocType).where(DocType.key == spec_4d.id))
        if doc_type:
            versions = (
                await s.scalars(
                    select(DocTypeVersion.id).where(DocTypeVersion.doc_type_id == doc_type.id)
                )
            ).all()
            await s.execute(delete(Document).where(Document.doc_type_version_id.in_(versions)))
            await s.execute(delete(DocType).where(DocType.id == doc_type.id))
            await s.commit()


# --------------------------------------------------------------------------- #
# bound to a version, carried forward
# --------------------------------------------------------------------------- #


async def test_a_new_version_inherits_the_template(published, spec_4d):
    """Re-uploading branding on every spec change is the tax that gets a feature
    abandoned. Publishing v2 keeps v1's template."""
    async with session() as s:
        v1 = await doc_types.get_version(s, published.id, 1)
        v1.template_uri = "s3://lcf-templates/fake/v1.docx"
        v1.template_filename = "branded.docx"
        await s.commit()

    spec_4d.version = 2
    async with session() as s:
        v2 = await doc_types.publish(s, spec_4d)
        await s.commit()
        assert v2.template_uri == "s3://lcf-templates/fake/v1.docx"
        assert v2.template_filename == "branded.docx"


async def test_a_carried_template_is_reported_as_incomplete_when_the_spec_grew(
    published, spec_4d, client
):
    """The case that makes carrying forward safe rather than silent: v2 adds a
    section, the inherited template predates it, and the card says so before a
    customer receives a document one chapter short."""
    async with session() as s:
        v1 = await doc_types.get_version(s, published.id, 1)
        original = docx_render.starter(doc_types.spec_of(v1))

    spec_4d.version = 2
    grown = spec_4d.model_copy(deep=True)
    grown.sections.append(
        Section(
            key="d9_lessons",
            title="D9 — Lessons Learned",
            blocks=[Block(key="lessons", kind=BlockKind.PROSE, label="Lessons learned")],
        )
    )

    async with session() as s:
        v2 = await doc_types.publish(s, grown)
        await s.commit()
        lint = templates_service.review(v2, original)

    assert lint.ok, "every tag it has is still real"
    assert not lint.complete
    assert any("d9_lessons" in m for m in lint.missing)


async def test_a_template_is_refused_when_it_names_what_the_version_lacks(published):
    async with session() as s:
        version = await doc_types.get_version(s, published.id, 1)
        bad = _template_with("{{ d9_nothing.text }}")
        with pytest.raises(templates_service.TemplateRejected) as raised:
            await templates_service.attach(s, version, bad, "bad.docx")
    assert any("no section 'd9_nothing'" in p for p in raised.value.lint.problems)


# --------------------------------------------------------------------------- #
# through HTTP — where the wiring is
# --------------------------------------------------------------------------- #


def test_the_starter_downloads_as_a_real_docx(published, client):
    response = client.get(f"/doc-types/{published.id}/versions/1/template/starter")
    assert response.status_code == 200
    assert response.content[:2] == b"PK"
    assert "template.docx" in response.headers["content-disposition"]

    doc = ReadDocx(io.BytesIO(response.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "{{ title }}" in text
    assert "D1 — Establish the Team" in text


def test_the_card_offers_the_starter_when_nothing_is_attached(published, client):
    response = client.get(f"/doc-types/{published.id}/versions/1/template/card")
    assert response.status_code == 200
    assert "Download starter template" in response.text
    assert "No template" in response.text


def test_the_version_page_loads_the_card(published, client):
    """A card reachable only by URL is a card nobody finds."""
    page = client.get(f"/doc-types/{published.id}/versions/1")
    assert f"/doc-types/{published.id}/versions/1/template/card" in page.text


def test_a_draft_does_not_offer_a_template(published, client, db):
    """A draft is not a version, so there is nothing to attach a template to."""
    started = client.post(f"/doc-types/{published.id}/build", follow_redirects=False)
    draft_id = started.headers["location"].rsplit("/", 1)[-1]
    read = client.get(f"/doc-types/drafts/{draft_id}/preview")
    assert read.status_code == 200
    assert "template/card" not in read.text


def test_uploading_a_bad_template_answers_with_the_card_and_attaches_nothing(published, client):
    """A service that refuses correctly and a form that never reaches it are
    indistinguishable from the service's side (ARCHITECTURE §15.4, point 10)."""
    bad = _template_with("{{ d9_nothing.text }}")
    response = client.post(
        f"/doc-types/{published.id}/versions/1/template",
        files={"file": ("bad.docx", bad, templates_service.DOCX_TYPE)},
    )
    assert response.status_code == 200
    assert "no section &#39;d9_nothing&#39;" in response.text or "d9_nothing" in response.text
    assert "attached</span>" not in response.text


def test_a_full_round_trip_brands_the_export(published, client, storage, spec_4d):
    """Download the starter, upload it back, export a document through it.

    The loop the feature exists for, walked the way a rule builder walks it.
    """
    starter = client.get(f"/doc-types/{published.id}/versions/1/template/starter").content

    attached = client.post(
        f"/doc-types/{published.id}/versions/1/template",
        files={"file": ("branded.docx", starter, templates_service.DOCX_TYPE)},
    )
    assert attached.status_code == 200
    assert "Template attached" in attached.text

    back = client.get(f"/doc-types/{published.id}/versions/1/template")
    assert back.content[:2] == b"PK", "what is in use must be readable back"

    created = client.post(
        "/documents",
        data={"doc_type": published.id, "title": "Templated"},
        follow_redirects=False,
    )
    document_id = created.headers["location"].rsplit("/", 1)[-1]

    export = client.post(
        f"/documents/{document_id}/export/docx",
        data={"override_reason": "rendering through the template under test"},
    )
    assert export.status_code == 200
    assert export.content[:2] == b"PK"
    # An empty document through the template still carries the template's shape.
    doc = ReadDocx(io.BytesIO(export.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "D1 — Establish the Team" in text
    assert "{{" not in text and "{%" not in text


def test_removing_a_template_falls_back_to_the_clean_layout(published, client, storage):
    starter = client.get(f"/doc-types/{published.id}/versions/1/template/starter").content
    client.post(
        f"/doc-types/{published.id}/versions/1/template",
        files={"file": ("branded.docx", starter, templates_service.DOCX_TYPE)},
    )
    removed = client.post(f"/doc-types/{published.id}/versions/1/template/remove")
    assert removed.status_code == 200
    assert "No template" in removed.text
    assert client.get(f"/doc-types/{published.id}/versions/1/template").status_code == 404


def _template_with(*lines: str) -> bytes:
    from docx import Document as NewDocx

    doc = NewDocx()
    for line in lines:
        doc.add_paragraph(line)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
