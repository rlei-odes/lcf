"""Export rendering.

All three formats come off one representation, so the tests check the same
document three ways: that nothing is invented, nothing filled in is lost, and the
structure comes from the spec rather than the content.
"""

import io
import json

import pytest
from docx import Document as ReadDocx
from docxtpl import DocxTemplate

from lcf.engine.view import DocumentView
from lcf.render import docx as docx_render
from lcf.render import markdown as markdown_render
from lcf.render import neutral


@pytest.fixture
def filled(spec_4d, sample_4d) -> DocumentView:
    content = {k: v.get("blocks", {}) for k, v in sample_4d["sections"].items()}
    answers = {k: v.get("answers", {}) for k, v in sample_4d["sections"].items()}
    return DocumentView(spec_4d, content, answers, completed=set(spec_4d.section_keys))


# --------------------------------------------------------------------------- #
# canonical JSON
# --------------------------------------------------------------------------- #


def test_json_carries_every_section_and_block(filled, spec_4d):
    payload = neutral.to_dict(filled, title="A 4D")
    assert [s["key"] for s in payload["sections"]] == spec_4d.section_keys

    header = next(s for s in payload["sections"] if s["key"] == "header")
    meta = next(b for b in header["blocks"] if b["key"] == "meta")
    assert meta["value"]["part_no"] == "A-4471"
    assert meta["kind"] == "keyvalue"


def test_json_records_the_doc_type_version(filled):
    payload = neutral.to_dict(filled, title="A 4D")
    assert payload["document"]["doc_type"] == "4d-report"
    assert payload["document"]["doc_type_version"] == 1
    assert payload["document"]["language"] == "en"


def test_json_is_serialisable(filled):
    """It is the handover format; it has to survive being written out."""
    text = json.dumps(neutral.to_dict(filled, title="A 4D"), ensure_ascii=False)
    assert json.loads(text)["document"]["title"] == "A 4D"


def test_json_reports_provenance_when_known(spec_4d):
    from datetime import UTC, datetime

    view = DocumentView(
        spec_4d,
        {"d2_problem": {"description": "text"}},
        provenance={
            "d2_problem": {
                "description": {
                    "author": "llm_accepted",
                    "actor": "local",
                    "at": datetime(2026, 9, 20, tzinfo=UTC),
                    "revisions": 2,
                }
            }
        },
    )
    payload = neutral.to_dict(view, title="x")
    block = next(
        b
        for s in payload["sections"]
        if s["key"] == "d2_problem"
        for b in s["blocks"]
        if b["key"] == "description"
    )
    assert block["provenance"]["author"] == "llm_accepted"
    assert block["provenance"]["revisions"] == 2
    assert block["provenance"]["at"].startswith("2026-09-20")


def test_empty_blocks_can_be_dropped(filled, spec_4d):
    view = DocumentView(spec_4d, {"header": {}})
    with_empty = neutral.to_dict(view, title="x", include_empty=True)
    without = neutral.to_dict(view, title="x", include_empty=False)
    assert sum(len(s["blocks"]) for s in with_empty["sections"]) > 0
    assert sum(len(s["blocks"]) for s in without["sections"]) == 0


# --------------------------------------------------------------------------- #
# markdown
# --------------------------------------------------------------------------- #


def test_markdown_structure_comes_from_the_spec(filled):
    text = markdown_render.render(filled, title="A 4D")
    assert text.startswith("# A 4D")
    assert "## D1 — Establish the Team" in text
    assert "### Team members" in text


def test_markdown_renders_a_table(filled):
    text = markdown_render.render(filled, title="A 4D")
    assert "| Name | Role | Function | Champion |" in text
    assert "| Sabine Vogt | 8D Champion | Quality management | yes |" in text
    assert "| Tomas Reiner |" in text


def test_markdown_renders_a_list(filled):
    text = markdown_render.render(filled, title="A 4D")
    assert "- Base material batch — certificates checked, within specification" in text


def test_markdown_skips_empty_sections_by_default(spec_4d):
    view = DocumentView(
        spec_4d,
        {
            "d1_team": {
                "members": [{"name": "A", "role": "B", "function": "C", "is_champion": True}]
            }
        },
    )
    text = markdown_render.render(view, title="x")
    assert "## D1 — Establish the Team" in text
    assert "## D4 — Root Cause Analysis" not in text


def test_markdown_cells_cannot_break_the_table(spec_4d):
    view = DocumentView(
        spec_4d,
        {
            "d1_team": {
                "members": [
                    {"name": "A | B", "role": "x\ny", "function": "f", "is_champion": False}
                ]
            }
        },
    )
    line = next(
        row for row in markdown_render.render(view, title="x").splitlines() if "A \\| B" in row
    )
    # The literal pipe is escaped, so the row still has one cell per column.
    cells = [c.strip() for c in line.replace("\\|", "\0").strip("|").split("|")]
    assert [c.replace("\0", "|") for c in cells] == ["A | B", "x y", "f", "no"]


def test_markdown_never_invents_content(spec_4d):
    """An empty document produces a title and nothing else."""
    text = markdown_render.render(DocumentView(spec_4d), title="Nothing yet")
    assert text.strip() == "# Nothing yet"


# --------------------------------------------------------------------------- #
# docx, no template
# --------------------------------------------------------------------------- #


def _docx_text(data: bytes) -> str:
    doc = ReadDocx(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_plain_docx_is_a_real_document(filled):
    data = docx_render.render_plain(filled, title="A 4D")
    assert data[:2] == b"PK", "a .docx is a zip"
    text = _docx_text(data)
    assert "A 4D" in text
    assert "D2 — Describe the Problem" in text
    assert "Nordwerk Fahrzeugtechnik" in text


def test_plain_docx_uses_real_tables(filled):
    doc = ReadDocx(io.BytesIO(docx_render.render_plain(filled, title="A 4D")))
    assert doc.tables, "tables must be Word tables, not pre-formatted text"
    headers = [c.text for c in doc.tables[0].rows[0].cells]
    assert "Report number" in headers or "Name" in headers


def test_plain_docx_omits_empty_sections(spec_4d):
    doc = ReadDocx(io.BytesIO(docx_render.render_plain(DocumentView(spec_4d), title="Empty")))
    assert [p.text for p in doc.paragraphs if p.text.strip()] == ["Empty"]


# --------------------------------------------------------------------------- #
# docx from a rule builder's template
# --------------------------------------------------------------------------- #


def _template(*lines: str) -> bytes:
    from docx import Document as NewDocx

    doc = NewDocx()
    for line in lines:
        doc.add_paragraph(line)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_a_template_placeholder_is_filled(filled):
    template = _template("Part: {{ header.meta }}", "Problem: {{ d2_problem.description }}")
    out = docx_render.render_with_template(filled, template, title="A 4D")
    text = _docx_text(out)
    assert "A-4471" in text
    assert "Nordwerk Fahrzeugtechnik reported cracking" in text


def test_template_context_exposes_section_titles(filled):
    context = docx_render.context(filled, title="A 4D")
    assert context["d1_team"]["title"] == "D1 — Establish the Team"
    assert "Sabine Vogt" in context["d1_team"]["members"]


def test_a_table_flattens_for_a_placeholder(filled):
    context = docx_render.context(filled, title="A 4D")
    assert "Name: Sabine Vogt" in context["d1_team"]["members"]


def test_a_rich_block_keeps_its_structure(filled):
    """A flattened table is a string; a branded template needs the rows."""
    ctx = docx_render.context(filled, title="A 4D")
    members = ctx["d1_team"]["members"]
    assert members.rows[0]["name"] == "Sabine Vogt"
    assert [c["key"] for c in members.columns][:2] == ["name", "role"]
    # The flat form still works, so a template written before this keeps working.
    assert "Name: Sabine Vogt" in members


def test_prose_becomes_paragraphs_not_asterisks(spec_4d):
    """Prose is markdown (DESIGN §9). `{{ }}` would print its syntax."""
    view = DocumentView(spec_4d, {"d2_problem": {"description": "A **bold** claim.\n\nSecond."}})
    value = docx_render.context(view, title="x")["d2_problem"]["description"]
    assert len(value.paragraphs) == 2
    assert "<w:b/>" in value.paragraphs[0].xml
    assert "**" not in value.paragraphs[0].xml


def test_an_empty_block_is_falsy_for_a_template(spec_4d):
    """`{% if section.block.filled %}` is how a template drops an empty chapter."""
    value = docx_render.context(DocumentView(spec_4d), title="x")["d2_problem"]["description"]
    assert not value.filled and value == ""


def test_linting_accepts_a_good_template(spec_4d):
    template = _template("{{ d2_problem.description }}", "{{ header.meta }}")
    assert docx_render.lint_template(spec_4d, template) == []


def test_linting_rejects_an_unknown_section(spec_4d):
    problems = docx_render.lint_template(spec_4d, _template("{{ d9_nothing.text }}"))
    assert problems and "no section 'd9_nothing'" in problems[0]


def test_linting_rejects_an_unknown_block(spec_4d):
    problems = docx_render.lint_template(spec_4d, _template("{{ d2_problem.invented }}"))
    assert problems and "has no block 'invented'" in problems[0]


def test_linting_allows_the_section_title(spec_4d):
    assert docx_render.lint_template(spec_4d, _template("{{ d2_problem.title }}")) == []


def test_linting_reads_tags_in_tables_and_headers(spec_4d):
    from docx import Document as NewDocx

    doc = NewDocx()
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "{{ d2_problem.ghost }}"
    doc.sections[0].header.paragraphs[0].text = "{{ nowhere.thing }}"
    buffer = io.BytesIO()
    doc.save(buffer)

    problems = docx_render.lint_template(spec_4d, buffer.getvalue())
    assert any("ghost" in p for p in problems), "a tag in a table cell must be linted"
    assert any("nowhere" in p for p in problems), "a tag in the header must be linted"


def test_a_linted_template_actually_renders(spec_4d, filled):
    """Linting clean must mean it works — otherwise the lint is theatre."""
    template = _template("{{ header.meta }}", "{{ d4_root_cause.causes }}")
    assert docx_render.lint_template(spec_4d, template) == []
    DocxTemplate(io.BytesIO(template))  # parses
    assert docx_render.render_with_template(filled, template, title="x")[:2] == b"PK"


# --------------------------------------------------------------------------- #
# the generated starter template
#
# The loop this has to keep true: a starter generated from a spec must lint clean
# against that spec and render that spec's content. If generation and linting ever
# disagree, the rule builder is handed a file the app then refuses.
# --------------------------------------------------------------------------- #


def _house_style() -> bytes:
    """Stand-in for a company .docx, until there is a UI to upload one.

    Everything a real one carries that the body does not: a header, a footer, a
    brand colour on a style, and placeholder prose that must not survive.
    """
    from docx import Document as NewDocx
    from docx.shared import RGBColor

    doc = NewDocx()
    doc.sections[0].header.paragraphs[0].text = "NORDWERK | Supplier Quality"
    doc.sections[0].footer.paragraphs[0].text = "Confidential"
    doc.styles["Heading 1"].font.color.rgb = RGBColor(0xC1, 0x00, 0x22)
    doc.styles["Heading 1"].font.name = "Georgia"
    doc.add_paragraph("Lorem ipsum from whoever made this template")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_a_generated_starter_lints_clean_against_its_own_spec(any_spec):
    """Over every example spec, not just the 4D: the generator has to be right
    for whatever a rule builder built, and this is the only test that would
    notice a block kind nobody wrote a tag for."""
    lint = docx_render.lint(any_spec, docx_render.starter(any_spec))
    assert lint.problems == []
    assert lint.missing == [], "the starter must mention every section and block"


def test_a_generated_starter_actually_renders_the_document(filled, spec_4d):
    out = docx_render.render_with_template(filled, docx_render.starter(spec_4d), title="A 4D")
    doc = ReadDocx(io.BytesIO(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]

    assert "Nordwerk Fahrzeugtechnik reported cracking" in text
    assert "Sabine Vogt" in cells, "a table row loop must produce real table rows"
    assert not any("{%" in c or "{{" in c for c in cells), "unrendered tag left in a table"
    assert "{%" not in text and "{{" not in text, "unrendered tag left in the body"


def test_table_rows_repeat_with_the_content(spec_4d, sample_4d):
    """The whole point of `{%tr %}`: the number of rows is the content's, not the
    template's. Three team members in, three rows out."""
    content = {k: v.get("blocks", {}) for k, v in sample_4d["sections"].items()}
    members = content["d1_team"]["members"]
    view = DocumentView(spec_4d, content, completed=set(spec_4d.section_keys))
    out = docx_render.render_with_template(view, docx_render.starter(spec_4d), title="x")

    doc = ReadDocx(io.BytesIO(out))
    team = next(t for t in doc.tables if t.rows[0].cells[0].text == "Name")
    assert len(team.rows) == len(members) + 1, "one header row plus one row per member"


def test_the_starter_is_built_onto_the_house_style(spec_4d):
    """Branding a template by hand every time it is regenerated is the tax that
    gets the feature abandoned."""
    doc = ReadDocx(io.BytesIO(docx_render.starter(spec_4d, base=_house_style())))
    assert doc.sections[0].header.paragraphs[0].text == "NORDWERK | Supplier Quality"
    assert doc.sections[0].footer.paragraphs[0].text == "Confidential"
    assert str(doc.styles["Heading 1"].font.color.rgb) == "C10022"
    assert "Lorem ipsum" not in "\n".join(p.text for p in doc.paragraphs)


def test_the_house_style_survives_all_the_way_to_the_export(filled, spec_4d):
    template = docx_render.starter(spec_4d, base=_house_style())
    doc = ReadDocx(io.BytesIO(docx_render.render_with_template(filled, template, title="A 4D")))
    assert doc.sections[0].header.paragraphs[0].text == "NORDWERK | Supplier Quality"
    assert str(doc.styles["Heading 1"].font.color.rgb) == "C10022"


def test_plain_render_can_use_the_house_style_too(filled):
    """A type with no template still goes out on company paper."""
    out = docx_render.render_plain(filled, title="A 4D", base=_house_style())
    doc = ReadDocx(io.BytesIO(out))
    assert doc.sections[0].header.paragraphs[0].text == "NORDWERK | Supplier Quality"
    assert "Nordwerk Fahrzeugtechnik reported cracking" in "\n".join(
        p.text for p in doc.paragraphs
    )


# --------------------------------------------------------------------------- #
# completeness: the check that matters once a type can grow
# --------------------------------------------------------------------------- #


def test_a_template_missing_a_section_is_reported(spec_4d):
    """The carried-forward case: a template written against v1, a v2 that added a
    section. It lints clean — every tag it has is real — and is a chapter short."""
    lint = docx_render.lint(spec_4d, _template("{{ d2_problem.description }}"))
    assert lint.problems == [], "nothing it names is wrong"
    assert not lint.complete
    assert any("D1" in m for m in lint.missing)
    assert any("Is / Is not" in m for m in lint.missing), "a missing block counts too"


def test_a_complete_template_reports_nothing_missing(spec_4d):
    assert docx_render.lint(spec_4d, docx_render.starter(spec_4d)).complete


def test_loop_variables_are_not_mistaken_for_sections(spec_4d):
    """`{{ row.name }}` inside a loop names a row, not a section. Reading it as
    one is how a linter ends up refusing its own generated template."""
    lint = docx_render.lint(
        spec_4d,
        _template("{%tr for row in d1_team.members.rows %}", "{{ row.name }}", "{%tr endfor %}"),
    )
    assert lint.problems == []


def test_a_loop_over_a_block_that_does_not_exist_is_caught(spec_4d):
    lint = docx_render.lint(spec_4d, _template("{%tr for row in d1_team.ghosts.rows %}"))
    assert any("has no block 'ghosts'" in p for p in lint.problems)


def test_an_attribute_the_block_kind_cannot_have_is_caught(spec_4d):
    """`.rows` on a prose block renders as nothing. Silently."""
    lint = docx_render.lint(spec_4d, _template("{%p for row in d2_problem.description.rows %}"))
    assert any("has no 'rows'" in p for p in lint.problems)
    assert any("paragraphs" in p for p in lint.problems), "say what it could have meant"
