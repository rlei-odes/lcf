"""Export: the gate rules, the record, and storage.

The rendering itself is covered in test_render.py without a database. What is
tested here is what may leave the building, and what gets written down when it
leaves anyway.
"""

import io
import json
import uuid

import pytest
from docx import Document as ReadDocx
from sqlalchemy import delete, select

from lcf.core.db import session
from lcf.models.tables import Assessment, DocType, DocTypeVersion, Document
from lcf.models.tables import CheckResult as CheckResultRow
from lcf.services import doc_types, documents, exports


@pytest.fixture
async def published(db, spec_4d):
    spec_4d.id = f"test-exp-{uuid.uuid4().hex[:8]}"
    async with session() as s:
        await doc_types.publish(s, spec_4d)
    yield spec_4d
    async with session() as s:
        doc_type = await s.scalar(select(DocType).where(DocType.key == spec_4d.id))
        if doc_type:
            versions = (
                await s.scalars(
                    select(DocTypeVersion.id).where(DocTypeVersion.doc_type_id == doc_type.id)
                )
            ).all()
            await s.execute(delete(Document).where(Document.doc_type_version_id.in_(versions)))
            await s.execute(delete(DocType).where(DocType.id == doc_type.id))


async def _filled_document(published, sample_4d, *, title="Export trial"):
    async with session() as s:
        document = await documents.create(s, published.id, title)
    for key, supplied in sample_4d["sections"].items():
        async with session() as s:
            if supplied.get("answers"):
                await documents.set_answers(s, document.id, key, supplied["answers"])
            for block_key, value in (supplied.get("blocks") or {}).items():
                await documents.set_block(s, document.id, key, block_key, value)
            await documents.mark_complete(s, document.id, key)
    return document


async def _pass_the_gate(document_id, spec):
    """Record passing judgements, as a real assessment run would."""
    from lcf.engine.checks.judged import pending_checks

    async with session() as s:
        assessment = Assessment(document_id=document_id, passed=True)
        s.add(assessment)
        await s.flush()
        for _kind, check, _section in pending_checks(spec):
            s.add(
                CheckResultRow(
                    assessment_id=assessment.id,
                    check_id=check.id,
                    species="judged",
                    result="pass",
                    severity=str(check.severity),
                    confidence=0.9,
                )
            )


# --------------------------------------------------------------------------- #
# what may leave
# --------------------------------------------------------------------------- #


async def test_an_unfinished_document_cannot_be_exported(published):
    async with session() as s:
        document = await documents.create(s, published.id, "Empty")
        with pytest.raises(exports.GateBlocked) as blocked:
            await exports.create(s, document.id, "docx")
    assert blocked.value.blockers > 0


async def test_unrun_checks_block_export_too(published, sample_4d):
    """Otherwise the gate is defeated by never running it."""
    document = await _filled_document(published, sample_4d)
    async with session() as s:
        with pytest.raises(exports.GateBlocked) as blocked:
            await exports.create(s, document.id, "markdown")

    assert blocked.value.blockers == 0, "the deterministic checks all pass"
    assert blocked.value.unchecked == 11
    assert "never run" in blocked.value.summary


async def test_a_passing_document_exports_without_ceremony(published, sample_4d, spec_4d):
    document = await _filled_document(published, sample_4d)
    await _pass_the_gate(document.id, spec_4d)

    async with session() as s:
        export, rendered = await exports.create(s, document.id, "markdown")

    assert export.gate_passed is True
    assert export.override_reason is None
    assert export.blockers == 0
    assert rendered.filename.endswith(".md")
    assert b"D4" in rendered.data


async def test_an_override_is_recorded_against_the_export(published, sample_4d):
    document = await _filled_document(published, sample_4d)
    async with session() as s:
        export, _ = await exports.create(
            s, document.id, "docx", override_reason="  customer deadline  "
        )

    assert export.gate_passed is False
    assert export.override_reason == "customer deadline", "trimmed, but kept"

    async with session() as s:
        past = await exports.history(s, document.id)
    assert [e.override_reason for e in past] == ["customer deadline"]


async def test_a_blank_override_is_not_an_override(published):
    async with session() as s:
        document = await documents.create(s, published.id, "Empty")
        with pytest.raises(exports.GateBlocked):
            await exports.create(s, document.id, "json", override_reason="   ")


# --------------------------------------------------------------------------- #
# the artefacts
# --------------------------------------------------------------------------- #


@pytest.mark.parametrize("fmt,head", [("json", b"{"), ("markdown", b"#"), ("docx", b"PK")])
async def test_every_format_renders(published, sample_4d, fmt, head):
    document = await _filled_document(published, sample_4d)
    async with session() as s:
        rendered = await exports.render(s, document.id, fmt)
    assert rendered.data[: len(head)] == head
    assert rendered.format == fmt


async def test_json_export_includes_the_assessment(published, sample_4d):
    document = await _filled_document(published, sample_4d)
    async with session() as s:
        rendered = await exports.render(s, document.id, "json")
    payload = json.loads(rendered.data)

    assert payload["document"]["id"] == str(document.id)
    assert payload["assessment"]["checks"], "the handover carries what was checked"
    assert any(c["result"] == "pass" for c in payload["assessment"]["checks"])


async def test_docx_export_contains_the_content(published, sample_4d):
    document = await _filled_document(published, sample_4d)
    async with session() as s:
        rendered = await exports.render(s, document.id, "docx")

    doc = ReadDocx(io.BytesIO(rendered.data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Export trial" in text
    assert "D3 — Interim Containment Actions" in text
    assert doc.tables, "the team and action tables are real Word tables"


async def test_an_unknown_format_is_refused(published):
    async with session() as s:
        document = await documents.create(s, published.id, "Empty")
        with pytest.raises(ValueError, match="unknown format"):
            await exports.render(s, document.id, "pdf")


async def test_the_filename_comes_from_the_title(published, sample_4d):
    document = await _filled_document(published, sample_4d, title="4D: bracket A-4471 / lot 8")
    async with session() as s:
        rendered = await exports.render(s, document.id, "docx")
    assert "/" not in rendered.filename
    assert ":" not in rendered.filename
    assert rendered.filename.endswith(".docx")


# --------------------------------------------------------------------------- #
# storage
# --------------------------------------------------------------------------- #


async def test_an_export_is_kept_and_can_be_fetched_back(published, sample_4d, spec_4d):
    document = await _filled_document(published, sample_4d)
    await _pass_the_gate(document.id, spec_4d)

    async with session() as s:
        export, rendered = await exports.create(s, document.id, "markdown")

    if export.uri is None:
        pytest.skip("object storage not configured")

    assert export.uri.startswith("s3://")
    assert exports.fetch(export) == rendered.data, "what comes back is what went in"
    assert export.size_bytes == len(rendered.data)
